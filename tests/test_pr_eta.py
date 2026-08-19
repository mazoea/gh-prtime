import sys
import os
import unittest
from datetime import datetime, timedelta
from unittest.mock import MagicMock

_this_dir = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(_this_dir, ".."))


def _build_week_year(created, closed):
    """
        Replicate the week iteration logic from process_one to test it.
    """
    end_d = closed or datetime.now()

    week_year = []
    d = created.date() if isinstance(created, datetime) else created
    d = d - timedelta(days=d.weekday())
    end_date = end_d.date() if isinstance(end_d, datetime) else end_d
    while d <= end_date:
        iso = d.isocalendar()
        week_year.append((iso[1], iso[0]))
        d += timedelta(weeks=1)
    return week_year


class Testbasic(unittest.TestCase):

    def testparse(self):
        """ testparse """
        from prtime import parse_eta_lines
        s = """
| Phases            | JH  |  JP  | TM |   JM | Total  |
|-----------------|----:|----:|-----:|-----:|-------:|
| ETA                  |  0  |    |     3 |      0 |        3 |
| Developing      |  5+4,5+8+8+5+8+5  |    0,5 + 1 |    0 |      0 |         45 |
| Review             |  4+4  |     2.5 + 0.5 + 1 + 1 |    0 |      0 |         13 |
| Total                |   -  |   -   |  -    |   -    |         61 |
| ETA est.             |      |       |       |         |     40  |
| ETA cust.           |   -  |   -  |   -   |   -     |        40 |
        """
        pr = MagicMock()
        pr.body = s
        lines, ignored = parse_eta_lines(pr)
        self.assertGreater(len(lines), 0)
        self.assertFalse(ignored)


class TestWeekIteration(unittest.TestCase):

    def test_same_week(self):
        """PR created and closed in the same week."""
        created = datetime(2025, 9, 15)  # Monday, week 38
        closed = datetime(2025, 9, 17)   # Wednesday, week 38
        weeks = _build_week_year(created, closed)
        self.assertEqual(weeks, [(38, 2025)])

    def test_two_weeks_same_year(self):
        """PR spanning two weeks in the same year."""
        created = datetime(2025, 9, 15)  # week 38
        closed = datetime(2025, 9, 22)   # week 39
        weeks = _build_week_year(created, closed)
        self.assertEqual(weeks, [(38, 2025), (39, 2025)])

    def test_cross_year_boundary(self):
        """PR crossing a year boundary."""
        created = datetime(2024, 12, 23)  # week 52 of 2024
        closed = datetime(2025, 1, 8)     # week 2 of 2025
        weeks = _build_week_year(created, closed)
        # Should include week 52/2024, week 1/2025, week 2/2025
        self.assertEqual(len(weeks), 3)
        self.assertEqual(weeks[0], (52, 2024))
        self.assertEqual(weeks[-1], (2, 2025))
        # No week 0 should exist
        week_nums = [w for w, y in weeks]
        self.assertNotIn(0, week_nums)

    def test_year_with_week_53(self):
        """Test a year that has ISO week 53 (e.g., 2020)."""
        created = datetime(2020, 12, 28)  # week 53 of 2020
        closed = datetime(2021, 1, 5)     # week 1 of 2021
        weeks = _build_week_year(created, closed)
        self.assertIn((53, 2020), weeks)
        self.assertIn((1, 2021), weeks)
        week_nums = [w for w, y in weeks]
        self.assertNotIn(0, week_nums)

    def test_long_running_pr(self):
        """PR open for many weeks does not produce infinite loop."""
        created = datetime(2025, 1, 6)   # week 2
        closed = datetime(2025, 9, 15)   # week 38
        weeks = _build_week_year(created, closed)
        # Should have ~37 weeks
        self.assertGreater(len(weeks), 30)
        self.assertLess(len(weeks), 40)
        # All week numbers should be valid (1-53)
        for w, y in weeks:
            self.assertGreaterEqual(w, 1)
            self.assertLessEqual(w, 53)

    def test_week_52_end(self):
        """PR ending in week 52 should not cause issues."""
        created = datetime(2025, 12, 15)  # week 51
        closed = datetime(2025, 12, 22)   # week 52
        weeks = _build_week_year(created, closed)
        self.assertEqual(len(weeks), 2)
        week_nums = [w for w, y in weeks]
        self.assertIn(51, week_nums)
        self.assertIn(52, week_nums)


class TestSumHours(unittest.TestCase):
    """
    `sum_hours` parses arithmetic out of PR-body table cells, which are
    fully attacker-controlled. These tests pin both directions:
      * legitimate arithmetic still evaluates to the right float;
      * anything that is not plain numeric arithmetic (function calls,
        names, attribute access, dunder/sandbox-escape tricks, operators
        outside +-*/) is rejected and returns -1.0 -- i.e. no code runs.
    """

    def setUp(self):
        from prtime import sum_hours
        self.sum_hours = sum_hours

    def test_valid_arithmetic_parses(self):
        cases = {
            "1": 1.0,
            "0": 0.0,
            "1+2": 3.0,
            "1.5*2": 3.0,
            "(1+0.25)*4": 5.0,
            "(1+2)*3": 9.0,
            "5+4+8": 17.0,
            "2.5 + 0.5 + 1 + 1": 5.0,
            "-2": -2.0,
            "10/4": 2.5,
        }
        for expr, expected in cases.items():
            with self.subTest(expr=expr):
                self.assertEqual(self.sum_hours(expr, "test"), expected)

    def test_non_arithmetic_rejected(self):
        # Empty / None / non-numeric / wrong characters -> -1.0
        for expr in (None, "", "   ", "abc", "1,2", "5,5", "1\n2",
                     "os.system", "1 ; 2", "1 & 2"):
            with self.subTest(expr=expr):
                self.assertEqual(self.sum_hours(expr, "test"), -1.0)

    def test_no_code_execution(self):
        # Classic eval()-injection payloads must never execute; each must
        # be rejected (return -1.0) either by the character allowlist or by
        # the AST node allowlist.
        payloads = [
            '__import__("os").system("echo PWNED")',
            'open("/etc/passwd").read()',
            '().__class__.__bases__[0].__subclasses__()',
            'eval("1")',
            'exec("x=1")',
            'globals()',
            'lambda: 1',
            '[x for x in range(3)]',
        ]
        for expr in payloads:
            with self.subTest(expr=expr):
                self.assertEqual(self.sum_hours(expr, "test"), -1.0)

    def test_ast_disallowed_operators_rejected(self):
        # These pass the character allowlist (only digits/./()/+-*/ ) but
        # produce AST nodes outside the numeric-arithmetic set, so the AST
        # walk must still reject them.
        for expr in ("2**10", "1//2"):
            with self.subTest(expr=expr):
                self.assertEqual(self.sum_hours(expr, "test"), -1.0)

    def test_overlong_input_rejected(self):
        from prtime import _HOURS_MAX_LEN
        too_long = "+".join(["1"] * (_HOURS_MAX_LEN))  # well over the cap
        self.assertGreater(len(too_long), _HOURS_MAX_LEN)
        self.assertEqual(self.sum_hours(too_long, "test"), -1.0)


class TestXlsxWriter(unittest.TestCase):
    """ --xlsx weekly-tab writer (no network). """

    def setUp(self):
        import prtime
        prtime.settings = {"devs": ["AY", "JH", "JS", "TM", "JM"]}
        prtime.hours_row.init()

    @staticmethod
    def _template_workbook(path):
        import openpyxl
        wb = openpyxl.Workbook()
        wb.active.title = "template"
        wb.save(path)

    def test_tab_name(self):
        from datetime import date
        from prtime import xlsx_tab_name
        self.assertEqual(xlsx_tab_name(date(2026, 8, 10)), "od 10-Aug-26")
        self.assertEqual(xlsx_tab_name(date(2026, 8, 3)), "od 3-Aug-26")

    def test_cell_coerce(self):
        from datetime import datetime, timezone
        from prtime import _xlsx_cell
        self.assertIsNone(_xlsx_cell(""))
        self.assertIsNone(_xlsx_cell(None))
        self.assertEqual(_xlsx_cell(23.5), 23.5)
        self.assertEqual(_xlsx_cell("advent"), "advent")
        # openpyxl rejects tz-aware datetimes -> must be stringified
        self.assertIsInstance(
            _xlsx_cell(datetime(2026, 8, 13, tzinfo=timezone.utc)), str)

    def test_write_rows_and_force(self):
        import os
        import tempfile
        import openpyxl
        from prtime import hours_row, write_rows_xlsx

        path = os.path.join(tempfile.mkdtemp(), "sheet.xlsx")
        self._template_workbook(path)

        r = hours_row()
        r[r.h_week] = 33
        r[r.h_customer] = "advent"
        r[r.h_issue] = "c-image-to-text/pull/2078:gIssue 1133"
        r[r.h_state] = "closed"
        r[r.h_phase_dev] = 23.5
        r.dev("TM", 23.5)

        n = write_rows_xlsx(path, "od 10-Aug-26", [r])
        self.assertEqual(n, 1)

        wb = openpyxl.load_workbook(path)
        self.assertIn("od 10-Aug-26", wb.sheetnames)
        ws = wb["od 10-Aug-26"]
        self.assertEqual(len(ws.defined_names), 10)      # sheet-scoped names re-created
        self.assertEqual(ws["A14"].value, 1)             # running #
        self.assertEqual(ws["B14"].value, 33)            # #Week
        self.assertEqual(ws["C14"].value, "advent")      # Cust
        self.assertEqual(ws["G14"].value, "closed")      # State
        self.assertEqual(ws["S14"].value, 23.5)          # Dev TM column

        # existing tab without --force must refuse
        with self.assertRaises(SystemExit):
            write_rows_xlsx(path, "od 10-Aug-26", [r])
        # --force overwrites
        self.assertEqual(write_rows_xlsx(path, "od 10-Aug-26", [r], force=True), 1)

    def test_row_cap_refused(self):
        import os
        import tempfile
        from prtime import (hours_row, write_rows_xlsx,
                            _XLSX_FIRST_ROW, _XLSX_LAST_ROW)

        path = os.path.join(tempfile.mkdtemp(), "sheet.xlsx")
        self._template_workbook(path)
        capacity = _XLSX_LAST_ROW - _XLSX_FIRST_ROW + 1
        too_many = [hours_row() for _ in range(capacity + 1)]
        with self.assertRaises(SystemExit):
            write_rows_xlsx(path, "od 5-Jan-26", too_many)


if __name__ == '__main__':
    # unittest.main()
    suite = unittest.TestSuite()
    suite.addTests(unittest.TestLoader().loadTestsFromTestCase(Testbasic))
    suite.addTests(unittest.TestLoader().loadTestsFromTestCase(TestWeekIteration))
    suite.addTests(unittest.TestLoader().loadTestsFromTestCase(TestSumHours))
    suite.addTests(unittest.TestLoader().loadTestsFromTestCase(TestXlsxWriter))
    unittest.TextTestRunner(verbosity=2).run(suite)
